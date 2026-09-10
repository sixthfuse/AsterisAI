"""Discover and enrich BCIT courses missed by the area course catalogues.

The original Asteris course extractor starts from BCIT's six area course
catalogues.  This companion tool uses official program matrices as an
independent discovery source, reconciles their course links against the live
database and canonical workbooks, and produces an audit before any import.

Dry-run audit (default):
    python bcit_course_gap_enrichment.py --seed-workbook "missing courses1.xlsx"

Apply only validated missing records to PostgreSQL:
    python bcit_course_gap_enrichment.py --seed-workbook "missing courses1.xlsx" --apply-db
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import time
from dataclasses import asdict, dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable
from urllib.parse import urljoin, urlparse

import openpyxl
import requests
from bs4 import BeautifulSoup, Tag
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

BASE_URL = "https://www.bcit.ca"
PROGRAM_CATALOGUES = {
    "Applied & Natural Sciences": f"{BASE_URL}/study/applied-natural-sciences/programs/",
    "Business & Media": f"{BASE_URL}/study/business-media/programs/",
    "Computing & IT": f"{BASE_URL}/study/computing-it/programs/",
    "Engineering": f"{BASE_URL}/study/engineering/programs/",
    "Health Sciences": f"{BASE_URL}/study/health-sciences/programs/",
    "Trades & Apprenticeships": f"{BASE_URL}/study/trades-apprenticeships/programs/",
}
CANONICAL_WORKBOOKS = (
    "BCIT_Courses_Phase1_2.xlsx",
    "BCIT_Courses_Phase1_2_Cleaned.xlsx",
    "BCIT_Courses_Phase1_2_Cleanedbycarlos.xlsx",
)
COURSE_PATH_RE = re.compile(r"^/courses/[^/?#]+/?$")
PROGRAM_PATH_RE = re.compile(r"^/programs/[^/?#]+/?$")
CODE_RE = re.compile(r"\b([A-Z]{3,5})\s*[- ]?\s*(\d{3,5}[A-Z]?)\b", re.I)
HEADERS = {
    "User-Agent": "Mozilla/5.0 (compatible; Asteris BCIT course-gap audit)",
    "Accept": "text/html,application/xhtml+xml",
    "Accept-Language": "en-CA,en;q=0.9",
}


@dataclass(frozen=True)
class Discovery:
    course_id: str
    display_course_code: str
    course_url: str
    discovery_source_url: str
    discovery_method: str


@dataclass
class CourseRecord:
    course_id: str
    display_course_code: str
    course_name: str
    credits: str
    course_overview_raw: str
    prerequisite_text_raw: str
    prerequisite_status: str
    course_url: str
    status: str
    availability_status: str
    page_title: str
    subject_category: str
    discovery_source_url: str
    discovery_method: str
    http_status: int
    validation_status: str
    validation_notes: str


def clean(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def normalize_code(value: str | None) -> tuple[str, str] | None:
    match = CODE_RE.search(clean(value).upper())
    if not match:
        return None
    subject, number = match.groups()
    return f"{subject}{number}", f"{subject} {number}"


def build_session() -> requests.Session:
    retry = Retry(
        total=5,
        connect=5,
        read=5,
        other=5,
        status=1,
        allowed_methods=frozenset({"GET"}),
        status_forcelist=(429, 500, 502, 503, 504),
        backoff_factor=1.0,
        respect_retry_after_header=True,
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retry, pool_connections=4, pool_maxsize=4)
    session = requests.Session()
    session.headers.update(HEADERS)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    return session


def get_soup(session: requests.Session, url: str, delay: float) -> tuple[BeautifulSoup, int]:
    response = session.get(url, timeout=(8, 15))
    status = response.status_code
    response.raise_for_status()
    if delay:
        time.sleep(delay)
    return BeautifulSoup(response.text, "html.parser"), status


def canonical_bcit_url(href: str, pattern: re.Pattern[str]) -> str | None:
    absolute = urljoin(BASE_URL, href)
    parsed = urlparse(absolute)
    if parsed.netloc.lower() not in {"bcit.ca", "www.bcit.ca"}:
        return None
    if not pattern.match(parsed.path):
        return None
    return f"{BASE_URL}{parsed.path.rstrip('/')}/"


def read_seed_codes(path: Path) -> list[str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    found: list[str] = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            continue
        code_columns = [
            index for index, value in enumerate(header)
            if clean(str(value or "")).lower().replace("_", " ") in {"course code", "course id"}
        ]
        for row in rows:
            for index in code_columns:
                if index >= len(row):
                    continue
                normalized = normalize_code(str(row[index] or ""))
                if normalized:
                    found.append(normalized[0])
    return sorted(set(found))


def workbook_course_ids(repo: Path) -> set[str]:
    course_ids: set[str] = set()
    for name in CANONICAL_WORKBOOKS:
        path = repo / name
        if not path.exists():
            continue
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                header = next(rows)
            except StopIteration:
                continue
            indexes = [
                index for index, value in enumerate(header)
                if clean(str(value or "")).lower() in {"course_code", "course code", "course_id"}
            ]
            for row in rows:
                for index in indexes:
                    if index < len(row):
                        normalized = normalize_code(str(row[index] or ""))
                        if normalized:
                            course_ids.add(normalized[0])
    return course_ids


def database_course_ids() -> set[str]:
    from database import get_connection

    with get_connection() as connection:
        with connection.cursor() as cursor:
            cursor.execute("SELECT course_id FROM courses")
            return {str(row[0]).upper() for row in cursor.fetchall()}


def discover_program_urls(session: requests.Session, delay: float) -> set[str]:
    urls: set[str] = set()
    for catalogue_url in PROGRAM_CATALOGUES.values():
        soup, _ = get_soup(session, catalogue_url, delay)
        for anchor in soup.find_all("a", href=True):
            url = canonical_bcit_url(anchor["href"], PROGRAM_PATH_RE)
            if url:
                urls.add(url)
    return urls


def course_discoveries_from_page(soup: BeautifulSoup, source_url: str) -> list[Discovery]:
    found: dict[tuple[str, str], Discovery] = {}
    for anchor in soup.find_all("a", href=True):
        url = canonical_bcit_url(anchor["href"], COURSE_PATH_RE)
        if not url:
            continue
        normalized = normalize_code(anchor.get_text(" ", strip=True))
        if not normalized:
            slug_match = re.search(r"-([a-z]{3,5})-(\d{3,5}[a-z]?)/$", url, re.I)
            normalized = normalize_code(" ".join(slug_match.groups())) if slug_match else None
        if not normalized:
            continue
        course_id, display = normalized
        found[(course_id, url)] = Discovery(course_id, display, url, source_url, "program_matrix_link")
    return list(found.values())


def discover_from_programs(
    session: requests.Session, program_urls: Iterable[str], delay: float
) -> list[Discovery]:
    discoveries: dict[tuple[str, str], Discovery] = {}
    for index, program_url in enumerate(sorted(set(program_urls)), 1):
        print(f"Program {index}: {program_url}")
        try:
            soup, _ = get_soup(session, program_url, delay)
        except requests.RequestException as error:
            print(f"  WARNING: {error}")
            continue
        for item in course_discoveries_from_page(soup, program_url):
            discoveries[(item.course_id, item.course_url)] = item
    return list(discoveries.values())


def seed_discoveries_from_programs(discoveries: Iterable[Discovery], seeds: Iterable[str]) -> dict[str, Discovery]:
    wanted = set(seeds)
    result: dict[str, Discovery] = {}
    for item in discoveries:
        if item.course_id in wanted:
            result.setdefault(item.course_id, item)
    return result


def section_text(soup: BeautifulSoup, heading_pattern: str, stop_headings: tuple[str, ...]) -> str:
    heading = soup.find(re.compile(r"^h[1-6]$"), string=re.compile(heading_pattern, re.I))
    if heading is None:
        for candidate in soup.find_all(re.compile(r"^h[1-6]$")):
            if re.search(heading_pattern, clean(candidate.get_text(" ", strip=True)), re.I):
                heading = candidate
                break
    if heading is None:
        return ""
    pieces: list[str] = []
    for sibling in heading.next_siblings:
        if not isinstance(sibling, Tag):
            continue
        if re.fullmatch(r"h[1-6]", sibling.name or ""):
            title = clean(sibling.get_text(" ", strip=True)).lower()
            if any(title.startswith(stop.lower()) for stop in stop_headings):
                break
        text = clean(sibling.get_text(" ", strip=True))
        if text:
            pieces.append(text)
    return clean(" ".join(pieces))


def page_status(soup: BeautifulSoup) -> str:
    text = clean(soup.get_text(" ", strip=True)).lower()
    if "this course is not offered this term" in text or "not offered this term" in text:
        return "Not offered this term"
    if "register now" in text or "add to cart" in text:
        return "Offered"
    return "Status not stated"


def parse_course_page(
    session: requests.Session, discovery: Discovery, delay: float
) -> CourseRecord:
    try:
        soup, http_status = get_soup(session, discovery.course_url, delay)
    except requests.RequestException as error:
        return CourseRecord(
            discovery.course_id, discovery.display_course_code, "", "", "", "",
            "unknown", discovery.course_url, "", "", "", "", discovery.discovery_source_url,
            discovery.discovery_method, getattr(error.response, "status_code", 0) or 0,
            "unresolved", str(error),
        )

    h1 = soup.find("h1")
    page_title = clean(h1.get_text(" ", strip=True) if h1 else "")
    normalized_title = normalize_code(page_title)
    title_without_code = clean(CODE_RE.sub("", page_title).strip(" ()-–—"))
    overview = section_text(soup, r"^Course Overview$", ("Prerequisite", "Credits", "Learning Outcomes"))
    prerequisites = section_text(soup, r"^Prerequisite\(s\)$|^Prerequisites$", ("Credits", "Learning Outcomes", "Related Programs"))
    credits_raw = section_text(soup, r"^Credits$", ("Learning Outcomes", "Related Programs", "Set Notifications"))
    credit_match = re.search(r"\b\d+(?:\.\d+)?\b", credits_raw)
    credits = credit_match.group(0) if credit_match else ""
    subtitle = h1.find_next(re.compile(r"^p$|^div$")) if h1 else None
    subject_category = clean(subtitle.get_text(" ", strip=True) if subtitle else "")

    notes: list[str] = []
    if not normalized_title or normalized_title[0] != discovery.course_id:
        notes.append("page heading course code does not match discovered code")
    if not title_without_code:
        notes.append("course name missing")
    if not overview:
        notes.append("course overview missing")
    if not credits:
        notes.append("credits missing")
    if not prerequisites:
        notes.append("prerequisite section missing")
    prereq_status = (
        "explicit_none" if re.search(r"no prerequisites? (?:are|is) required", prerequisites, re.I)
        else "authoritative_raw_text" if prerequisites
        else "unknown"
    )
    validation = "validated" if not notes else "needs_review"
    return CourseRecord(
        course_id=discovery.course_id,
        display_course_code=normalized_title[1] if normalized_title else discovery.display_course_code,
        course_name=title_without_code,
        credits=credits,
        course_overview_raw=overview,
        prerequisite_text_raw=prerequisites,
        prerequisite_status=prereq_status,
        course_url=discovery.course_url,
        status="Active",
        availability_status=page_status(soup),
        page_title=page_title,
        subject_category=subject_category,
        discovery_source_url=discovery.discovery_source_url,
        discovery_method=discovery.discovery_method,
        http_status=http_status,
        validation_status=validation,
        validation_notes="; ".join(notes),
    )


def write_audit(output_dir: Path, rows: list[dict], summary: dict) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "bcit_course_gap_audit.json"
    csv_path = output_dir / "bcit_course_gap_audit.csv"
    json_path.write_text(json.dumps({"summary": summary, "courses": rows}, indent=2, ensure_ascii=False), encoding="utf-8")
    fieldnames = sorted({key for row in rows for key in row})
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def apply_database(records: Iterable[CourseRecord]) -> int:
    from database import get_connection

    inserted = 0
    sql = """
        INSERT INTO courses (
            course_id, course_name, credits, course_overview, status, source_url,
            last_checked, notes, display_course_code, area_of_study,
            subject_category, source_catalogue_url, source_course_url, page_title,
            course_overview_raw, prerequisite_text_raw, prerequisite_text_clean,
            description_status, prerequisite_status, import_source
        ) VALUES (
            %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
            %s, %s, %s, %s, %s, %s, %s, %s, %s, %s
        ) ON CONFLICT (course_id) DO NOTHING
    """
    with get_connection() as connection:
        with connection.cursor() as cursor:
            for record in records:
                if record.validation_status != "validated":
                    continue
                try:
                    credits = Decimal(record.credits)
                except InvalidOperation:
                    continue
                cursor.execute(sql, (
                    record.course_id, record.course_name, credits,
                    record.course_overview_raw, record.status, record.course_url,
                    date.today(), record.prerequisite_text_raw,
                    record.display_course_code, "", record.subject_category,
                    record.discovery_source_url, record.course_url, record.page_title,
                    record.course_overview_raw, record.prerequisite_text_raw,
                    record.prerequisite_text_raw, "source_preserved",
                    record.prerequisite_status, "bcit_program_matrix_gap_enrichment",
                ))
                inserted += cursor.rowcount
        connection.commit()
    return inserted


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--seed-workbook", type=Path, required=True)
    parser.add_argument("--repo", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent / "course_gap_audit")
    parser.add_argument("--program-url", action="append", default=[], help="Audit only these official program pages; repeatable")
    parser.add_argument("--all-programs", action="store_true", help="Discover and scan program pages from all six official catalogues")
    parser.add_argument(
        "--enrich-seeds-only",
        action="store_true",
        help="For a broad discovery audit, enrich only seed courses and leave other gaps as discovered candidates",
    )
    parser.add_argument("--delay", type=float, default=0.15)
    parser.add_argument("--apply-db", action="store_true")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    session = build_session()
    seeds = read_seed_codes(args.seed_workbook)
    db_ids = database_course_ids()
    workbook_ids = workbook_course_ids(args.repo)

    program_urls = set(args.program_url)
    if args.all_programs:
        program_urls.update(discover_program_urls(session, args.delay))
    if not program_urls:
        program_urls.add("https://www.bcit.ca/programs/electronics-bachelor-of-technology-part-time-8900btech/")

    discoveries = discover_from_programs(session, program_urls, args.delay)
    seed_map = seed_discoveries_from_programs(discoveries, seeds)
    targets: dict[str, Discovery] = dict(seed_map)
    if args.all_programs:
        for item in discoveries:
            if item.course_id not in db_ids and item.course_id not in workbook_ids:
                targets.setdefault(item.course_id, item)

    enrichment_targets = (
        {course_id: item for course_id, item in targets.items() if course_id in set(seeds)}
        if args.enrich_seeds_only else targets
    )
    records = [
        parse_course_page(session, enrichment_targets[course_id], args.delay)
        for course_id in sorted(enrichment_targets)
    ]
    rows: list[dict] = []
    for record in records:
        row = asdict(record)
        row["seed"] = record.course_id in seeds
        row["in_database_before"] = record.course_id in db_ids
        row["in_canonical_workbooks_before"] = record.course_id in workbook_ids
        row["reconciliation"] = (
            "existing_database" if row["in_database_before"]
            else "existing_workbook" if row["in_canonical_workbooks_before"]
            else "newly_found" if record.validation_status == "validated"
            else "unresolved"
        )
        rows.append(row)

    enriched_ids = {record.course_id for record in records}
    for course_id, discovery in sorted(targets.items()):
        if course_id in enriched_ids:
            continue
        row = asdict(discovery)
        row.update({
            "seed": course_id in seeds,
            "in_database_before": course_id in db_ids,
            "in_canonical_workbooks_before": course_id in workbook_ids,
            "reconciliation": "discovered_unenriched",
            "validation_status": "official_program_link_candidate",
            "validation_notes": "Official BCIT program page links to this course page; course detail page not enriched in this run.",
        })
        rows.append(row)

    missing_seed_links = sorted(set(seeds) - set(seed_map))
    summary = {
        "run_date": date.today().isoformat(),
        "seed_count": len(seeds),
        "program_pages_scanned": len(program_urls),
        "unique_program_course_links": len({item.course_id for item in discoveries}),
        "database_count_before": len(db_ids),
        "canonical_workbook_unique_codes": len(workbook_ids),
        "existing_database": sum(row["reconciliation"] == "existing_database" for row in rows),
        "existing_workbook_only": sum(row["reconciliation"] == "existing_workbook" for row in rows),
        "newly_found_validated": sum(row["reconciliation"] == "newly_found" for row in rows),
        "additional_official_program_link_candidates": sum(
            row["reconciliation"] == "discovered_unenriched" and not row["seed"] for row in rows
        ),
        "unresolved_or_ambiguous": sum(row["reconciliation"] == "unresolved" for row in rows) + len(missing_seed_links),
        "missing_seed_links": missing_seed_links,
        "database_inserted": 0,
    }
    if args.apply_db:
        summary["database_inserted"] = apply_database(records)
        summary["database_count_after"] = len(database_course_ids())
    write_audit(args.output_dir, rows, summary)
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
